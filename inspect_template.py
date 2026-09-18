"""Inspect the template PPTX to understand available layouts and placeholders."""
from pptx import Presentation
from pptx.util import Inches, Pt

template_path = r"c:\Users\jasil_myg\Desktop\OSG-myG-PORTAL-mainnnnn - Copy\Challenging Success Business PowerPoint Templates.pptx"
prs = Presentation(template_path)

print(f"Slide size: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")
print(f"Total slide layouts: {len(prs.slide_layouts)}")
print()

for i, layout in enumerate(prs.slide_layouts):
    print(f"Layout {i}: '{layout.name}'")
    for ph in layout.placeholders:
        print(f"    PH idx={ph.placeholder_format.idx}  type={ph.placeholder_format.type}  name='{ph.name}'")

print("\n--- Existing slides in template ---")
for i, slide in enumerate(prs.slides):
    print(f"Slide {i}: layout='{slide.slide_layout.name}'  shapes={len(slide.shapes)}")
    for ph in slide.placeholders:
        try:
            txt = ph.text[:60].replace('\n', ' ')
        except Exception:
            txt = "[no text]"
        print(f"    PH idx={ph.placeholder_format.idx}  '{txt}'")
