"""
Q2 2026 Professional Work Summary — PowerPoint Generator
Business Template with Dark Navy + Gold Accent Theme
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Theme Colors ──────────────────────────────────────────────
NAVY      = RGBColor(0x0B, 0x1D, 0x3A)   # Deep navy background
DARK_NAVY = RGBColor(0x06, 0x12, 0x28)   # Darker navy
GOLD      = RGBColor(0xD4, 0xA5, 0x37)   # Gold accent
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT     = RGBColor(0xCC, 0xD6, 0xE0)   # Light gray-blue text
SOFT_GOLD = RGBColor(0xF0, 0xD0, 0x78)   # Softer gold
TEAL      = RGBColor(0x00, 0xB4, 0xD8)   # Teal accent
DARK_GRAY = RGBColor(0x1A, 0x2A, 0x40)   # Card background
MED_GRAY  = RGBColor(0x8B, 0x95, 0xA5)   # Subtle text
GREEN     = RGBColor(0x2E, 0xCC, 0x71)   # Success green
SLATE     = RGBColor(0x14, 0x22, 0x38)   # Slide bg variant

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


# ══════════════════════════════════════════════════════════════
# Helper Functions
# ══════════════════════════════════════════════════════════════

def add_bg(slide, color=NAVY):
    """Fill slide background with solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, color, alpha=None):
    """Add a filled rectangle shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_rounded_rect(slide, left, top, width, height, color):
    """Add a rounded rectangle card."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box with single-run formatting."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_gold_line(slide, left, top, width, height=Inches(0.04)):
    """Add a thin gold accent line."""
    return add_shape(slide, left, top, width, height, GOLD)

def add_bullet_slide_content(slide, items, start_top, left=Inches(0.9), width=Inches(11.5),
                              font_size=14, color=LIGHT, line_spacing=Pt(22)):
    """Add bulleted items to a slide."""
    txBox = slide.shapes.add_textbox(left, start_top, width, Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(6)
        p.space_before = Pt(2)
    return txBox

def section_header(slide, title, subtitle="", top=Inches(0.4)):
    """Add a section title with gold underline."""
    add_text_box(slide, Inches(0.8), top, Inches(11), Inches(0.6),
                 title, font_size=30, color=GOLD, bold=True, font_name="Calibri Light")
    add_gold_line(slide, Inches(0.8), top + Inches(0.55), Inches(2.5))
    if subtitle:
        add_text_box(slide, Inches(0.8), top + Inches(0.65), Inches(11), Inches(0.5),
                     subtitle, font_size=14, color=MED_GRAY, font_name="Calibri")

def add_kpi_card(slide, left, top, value, label, accent_color=GOLD):
    """Add a KPI metric card."""
    card = add_rounded_rect(slide, left, top, Inches(2.6), Inches(1.4), DARK_GRAY)
    # Accent line at top of card
    add_shape(slide, left + Inches(0.05), top + Inches(0.05), Inches(2.5), Inches(0.05), accent_color)
    # Value
    add_text_box(slide, left + Inches(0.15), top + Inches(0.2), Inches(2.3), Inches(0.6),
                 value, font_size=28, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Label
    add_text_box(slide, left + Inches(0.15), top + Inches(0.8), Inches(2.3), Inches(0.5),
                 label, font_size=11, color=MED_GRAY, alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# SLIDE 1: Title Slide
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide, DARK_NAVY)

# Left gold accent bar
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, GOLD)

# Top subtle line
add_shape(slide, Inches(0.08), Inches(0), W, Inches(0.02), GOLD)

# Bottom subtle line
add_shape(slide, Inches(0.08), H - Inches(0.02), W, Inches(0.02), GOLD)

# Title block
add_text_box(slide, Inches(1.2), Inches(1.5), Inches(10), Inches(0.5),
             "DETAILED WORK SUMMARY", font_size=18, color=GOLD, bold=True,
             font_name="Calibri", alignment=PP_ALIGN.LEFT)

add_text_box(slide, Inches(1.2), Inches(2.1), Inches(10), Inches(1.2),
             "Q2 2026: April | May | June", font_size=44, color=WHITE, bold=True,
             font_name="Calibri Light", alignment=PP_ALIGN.LEFT)

add_gold_line(slide, Inches(1.2), Inches(3.4), Inches(4), Inches(0.05))

add_text_box(slide, Inches(1.2), Inches(3.7), Inches(10), Inches(0.6),
             "Comprehensive breakdown of engineering, database optimization,\nfrontend development, and AI integration across all workspaces",
             font_size=16, color=LIGHT, font_name="Calibri")

add_text_box(slide, Inches(1.2), Inches(5.5), Inches(5), Inches(0.4),
             "Jasil N Balussery  ·  Technology / Software Development  ·  myG",
             font_size=13, color=MED_GRAY, font_name="Calibri")

add_text_box(slide, Inches(1.2), Inches(5.9), Inches(5), Inches(0.4),
             "Reporting Period: 01 April 2026 – 30 June 2026",
             font_size=12, color=MED_GRAY, font_name="Calibri")

# Right side decorative element
add_shape(slide, Inches(10.5), Inches(1.0), Inches(0.03), Inches(5.5), GOLD)
add_text_box(slide, Inches(10.8), Inches(2.5), Inches(2), Inches(0.5),
             "8", font_size=72, color=GOLD, bold=True, alignment=PP_ALIGN.CENTER, font_name="Calibri Light")
add_text_box(slide, Inches(10.8), Inches(3.8), Inches(2), Inches(0.5),
             "PROJECTS\nDELIVERED", font_size=14, color=MED_GRAY, alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# SLIDE 2: Executive Summary & KPIs
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "Executive Summary & Key Metrics")

# Summary text
add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.2),
             "During Q2 2026, extensive full-stack development, database engineering, AI/ML integration, "
             "and multiple greenfield product launches were delivered across 8 distinct projects. "
             "Work spanned Python (Flask/Django), PostgreSQL, JavaScript, Machine Learning, "
             "LLM integration (NVIDIA Nemotron/OpenRouter), Google APIs, and cloud deployment (Render/Vercel).",
             font_size=14, color=LIGHT)

# KPI Cards Row 1
add_kpi_card(slide, Inches(0.6),  Inches(2.6), "8",       "Projects Delivered", GOLD)
add_kpi_card(slide, Inches(3.5),  Inches(2.6), "80+",     "Dev Conversations", TEAL)
add_kpi_card(slide, Inches(6.4),  Inches(2.6), "12.6M+",  "Database Rows Managed", GREEN)
add_kpi_card(slide, Inches(9.3),  Inches(2.6), "26",      "Materialized Views", GOLD)

# KPI Cards Row 2
add_kpi_card(slide, Inches(0.6),  Inches(4.3), "6",       "APIs Integrated", TEAL)
add_kpi_card(slide, Inches(3.5),  Inches(4.3), "3",       "Cloud Platforms", GREEN)
add_kpi_card(slide, Inches(6.4),  Inches(4.3), "4",       "ML Models Deployed", GOLD)
add_kpi_card(slide, Inches(9.3),  Inches(4.3), "< 1s",   "AI Response Time\n(from 5-10 min)", TEAL)

# Tech stack bar at bottom
add_shape(slide, Inches(0.6), Inches(6.1), Inches(11.7), Inches(1.0), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(6.15), Inches(2), Inches(0.3),
             "TECH STACK", font_size=10, color=GOLD, bold=True)
add_text_box(slide, Inches(0.8), Inches(6.45), Inches(11.3), Inches(0.6),
             "Python  ·  Flask  ·  Django  ·  PostgreSQL  ·  Redis  ·  JavaScript  ·  Scikit-Learn  ·  LSTM  ·  "
             "NVIDIA NIM  ·  OpenRouter  ·  n8n  ·  Telinfy API  ·  Google Sheets/Drive API  ·  "
             "Plotly  ·  Chart.js  ·  Render  ·  Vercel  ·  Google OAuth 2.0",
             font_size=11, color=LIGHT)


# ══════════════════════════════════════════════════════════════
# SLIDE 3: OSG-myG-PORTAL (Claims & WhatsApp)
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "1. OSG-myG-PORTAL", "Claims Management & WhatsApp Automation")

# Left column - WhatsApp & Claims
card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(5.8), Inches(2.4), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(5), Inches(0.4),
             "⚡ WhatsApp API Integration (Telinfy/GreenAds)", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Engineered automated messaging pipelines linking portal events to WhatsApp API endpoints",
    "• Debugged payload routing and delivery issues for reliable outbound notifications",
    "• Configured 4 WhatsApp templates: Registered, Replacement, Repair Completed, Spare Parts",
    "• Built trigger buttons in portal for WhatsApp message dispatch",
], Inches(1.95), left=Inches(0.9), width=Inches(5.3), font_size=12)

card2 = add_rounded_rect(slide, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.4), DARK_GRAY)
add_text_box(slide, Inches(7.0), Inches(1.55), Inches(5), Inches(0.4),
             "🔍 Claims Processing & Data Forensics", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Developed investigative scripts to trace and resolve claim anomalies",
    "• Built repair utilities to retroactively fix corrupted claim states",
    "• SR Number auto-generation on status change (Submit → Registered)",
    "• Google Sheets bi-directional sync for all claim operations",
], Inches(1.95), left=Inches(7.1), width=Inches(5.3), font_size=12)

# Bottom row
card3 = add_rounded_rect(slide, Inches(0.6), Inches(4.2), Inches(5.8), Inches(2.4), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(4.25), Inches(5), Inches(0.4),
             "📊 Large-Scale Data Ingestion (OSID)", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Engineered massive data parsers for legacy Excel imports (17MB+ OSID files)",
    "• Automated data sanitization: European/American date standardization",
    "• Built comprehensive validation layer (15,493 bytes of validation logic)",
    "• Updated OSID reference data, Future Store List, RBM/BDM/Branch hierarchy",
], Inches(4.65), left=Inches(0.9), width=Inches(5.3), font_size=12)

card4 = add_rounded_rect(slide, Inches(6.8), Inches(4.2), Inches(5.8), Inches(2.4), DARK_GRAY)
add_text_box(slide, Inches(7.0), Inches(4.25), Inches(5), Inches(0.4),
             "🔗 Google Apps Script Bridge", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Developed seamless sync between Google Sheets and backend",
    "• New claims auto-append to Sheet with full data columns",
    "• Remarks & status updates sync to Follow-up Notes in real-time",
    "• Multiple Apps Script Web App deployments for reliable sync",
], Inches(4.65), left=Inches(7.1), width=Inches(5.3), font_size=12)


# ══════════════════════════════════════════════════════════════
# SLIDE 4: Enterprise AI Agent
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "2. Enterprise AI Agent", "Loyalty Portal — Natural Language Business Intelligence")

# Architecture flow
add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.6), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(3), Inches(0.3),
             "AI QUERY PIPELINE", font_size=11, color=GOLD, bold=True)

# Flow boxes
flow_items = [
    ("User Question", TEAL),
    ("→", WHITE),
    ("Schema Catalog", GOLD),
    ("→", WHITE),
    ("AI SQL Agent", TEAL),
    ("→", WHITE),
    ("PostgreSQL Engine", GOLD),
    ("→", WHITE),
    ("AI Analyst", TEAL),
    ("→", WHITE),
    ("Business Insight", GREEN),
]
x = Inches(0.8)
for text, color in flow_items:
    if text == "→":
        add_text_box(slide, x, Inches(2.05), Inches(0.35), Inches(0.4),
                     "→", font_size=20, color=GOLD, alignment=PP_ALIGN.CENTER)
        x += Inches(0.35)
    else:
        box = add_rounded_rect(slide, x, Inches(1.95), Inches(1.3), Inches(0.5), SLATE)
        add_text_box(slide, x, Inches(2.0), Inches(1.3), Inches(0.4),
                     text, font_size=9, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        x += Inches(1.35)

# Details cards
card1 = add_rounded_rect(slide, Inches(0.6), Inches(3.4), Inches(5.8), Inches(3.5), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(3.45), Inches(5), Inches(0.4),
             "🧠 Multi-Layered Agent Architecture", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• SQL Agent: Translates natural language → complex PostgreSQL queries",
    "• AI Analyst: Interprets raw DB output → readable business insights",
    "• FastPath Engine: Common queries answered in < 0.2 seconds",
    "• 16-phase implementation: Foundation → Security → Forecasting → Export",
    "• Schema context injection prevents hallucinations",
    "• Conversation memory for contextual follow-up questions",
], Inches(3.85), left=Inches(0.9), width=Inches(5.3), font_size=12)

card2 = add_rounded_rect(slide, Inches(6.8), Inches(3.4), Inches(5.8), Inches(3.5), DARK_GRAY)
add_text_box(slide, Inches(7.0), Inches(3.45), Inches(5), Inches(0.4),
             "⚡ LLM Integration & Optimization", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• NVIDIA NIM API (Nemotron Ultra) for SQL generation",
    "• OpenRouter API as fallback LLM provider",
    "• Response time reduced: 5-10 minutes → < 1 second",
    "• Aggressive prompt engineering & query scope reduction",
    "• n8n workflow automation integration for chat widget",
    "• Validated with 20+ test business questions",
], Inches(3.85), left=Inches(7.1), width=Inches(5.3), font_size=12)


# ══════════════════════════════════════════════════════════════
# SLIDE 5: Database Administration
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "3. Database Administration & Optimization", "12.6M+ Rows · DigitalOcean PostgreSQL")

# Three column cards
card_w = Inches(3.7)
gap = Inches(0.35)

card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), card_w, Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(3.3), Inches(0.4),
             "📐 Materialized View Overhaul", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Diagnosed critical cross-join bug in mv_yearly_cohort",
    "• Bug was duplicating rows and inflating LTV metrics by 1000%+",
    "• Built robust refresh system for 26 materialized views",
    "• Concurrent regeneration without database locks",
    "• Views include: mv_action_engine, mv_cohort_customer_years, mv_monthly_summary, mv_rfm_summary, mv_dormant_reactivation, and 21 more",
], Inches(1.95), left=Inches(0.8), width=Inches(3.3), font_size=11)

card2 = add_rounded_rect(slide, Inches(0.6) + card_w + gap, Inches(1.5), card_w, Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(0.6) + card_w + gap + Inches(0.2), Inches(1.55), Inches(3.3), Inches(0.4),
             "🚀 Caching Infrastructure", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Integrated Redis for production caching",
    "• LocMemCache for local development",
    "• Heavy API responses cached → millisecond load times",
    "• Fixed stale memory cache issues in local environments",
    "• 12.6M+ row calculations reduced to sub-second responses",
], Inches(1.95), left=Inches(0.6) + card_w + gap + Inches(0.2), width=Inches(3.3), font_size=11)

card3 = add_rounded_rect(slide, Inches(0.6) + 2*(card_w + gap), Inches(1.5), card_w, Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(0.6) + 2*(card_w + gap) + Inches(0.2), Inches(1.55), Inches(3.3), Inches(0.4),
             "📥 DB Manager & Data Ops", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Custom scripts bypass web-server timeouts for large uploads",
    "• DSR MAY & JUNE 2026 data imported (multiple parts)",
    "• Automated scrubbing of anomalous data (SMC/EI, HEAD OFFICE, UG SMART CHOICE)",
    "• Data integrity verification: 50,33,297+ total records",
    "• Upload success/failure notifications added",
], Inches(1.95), left=Inches(0.6) + 2*(card_w + gap) + Inches(0.2), width=Inches(3.3), font_size=11)


# ══════════════════════════════════════════════════════════════
# SLIDE 6: Machine Learning & Forecasting
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "4. Machine Learning & Predictive Forecasting",
               "Live Forecasting Engines Replacing Static Placeholders")

card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(5.8), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(5), Inches(0.4),
             "🤖 Model Deployment (Scikit-Learn)", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Random Forest — Ensemble decision tree model for sales prediction",
    "• MLPRegressor — Neural Network proxy for pattern recognition",
    "• GradientBoostingRegressor — Sequential boosting for accuracy",
    "• All models deployed directly into CampaignAnalysisAPIView",
    "• LSTM deep learning model trained on 2020-2026 historical data",
], Inches(1.95), left=Inches(0.9), width=Inches(5.3), font_size=12)

card2 = add_rounded_rect(slide, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(7.0), Inches(1.55), Inches(5), Inches(0.4),
             "🎯 Feature Engineering", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• MalayalamCalendarFeaturizer — Custom Kerala-specific feature engine",
    "• Converts dates → multi-dimensional arrays with festival proximity",
    "• Factors: Onam, Vishu, seasonal patterns, year-over-year trends",
    "• Weather factor integration for Kerala-specific predictions",
    "• Real-time sales prediction from mid-day actuals to EOD forecast",
], Inches(1.95), left=Inches(7.1), width=Inches(5.3), font_size=12)

card3 = add_rounded_rect(slide, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(4.35), Inches(5), Inches(0.4),
             "💤 Dormant Customer Reactivation System", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• SQL pre-aggregation engines (mv_dormant_reactivation) bucket customers: purchased 2024, silent 2025/2026",
    "• Cohort-based analysis across 2020-2025 cohorts with Resurrection Rate % tracking",
    "• UI Visualizers: Plotly Probability Gauges, Semi-Donut Charts, Dormancy Risk Meters",
    "• Features: Hidden patterns, seasonal comeback trends, festival reactivation spikes, repeat purchase probability",
    "• Branch-level drill-down: New vs Repeat customer analysis for specific Future stores",
], Inches(4.75), left=Inches(0.9), width=Inches(11.5), font_size=12)


# ══════════════════════════════════════════════════════════════
# SLIDE 7: SHE START Dashboard
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "5. SHE START — Applicant Evaluation Dashboard",
               "\"She Start — Her Dreams Start Here\" Startup Program")

card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(3.7), Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(3.3), Inches(0.4),
             "🔄 Live Google Sheets Sync", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• gspread library for bidirectional sync",
    "• 25-second auto-refresh interval",
    "• Data order matches Google Sheet exactly",
    "• Auto-detection of new applicants",
    "• District & business name auto-population",
    "• 28 candidates tracked",
], Inches(1.95), left=Inches(0.8), width=Inches(3.3), font_size=11)

card2 = add_rounded_rect(slide, Inches(4.65), Inches(1.5), Inches(3.7), Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(4.85), Inches(1.55), Inches(3.3), Inches(0.4),
             "📊 6-Panelist Scoring Algorithm", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• 6 panelists score each applicant",
    "• Auto-drop highest & lowest scores",
    "• Average remaining 4 for unbiased result",
    "• Weighted Final Score:",
    "  — Interview Score: 40%",
    "  — Growth Potential: 15%",
    "  — Support Need: 15%",
    "  — Emotional Impact: 10%",
    "  — Sustainability: 10%",
    "  — Utilization: 10%",
], Inches(1.95), left=Inches(4.85), width=Inches(3.3), font_size=11)

card3 = add_rounded_rect(slide, Inches(8.7), Inches(1.5), Inches(3.7), Inches(5.2), DARK_GRAY)
add_text_box(slide, Inches(8.9), Inches(1.55), Inches(3.3), Inches(0.4),
             "🖥️ Interactive Dashboard", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Inline cell editing for 5 score columns",
    "• Silent auto-save (no page reload)",
    "• Auto badging: Strong Selection (85+), Waitlist, etc.",
    "• Tab navigation fixed (→ direction)",
    "• Interview columns: non-editable",
    "• Dedicated user: shestart / shestart123",
    "• Excel report download",
    "• Deployed on Render",
], Inches(1.95), left=Inches(8.9), width=Inches(3.3), font_size=11)


# ══════════════════════════════════════════════════════════════
# SLIDE 8: Enterprise Retail Dashboard
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "6. Enterprise Retail Dashboard Enhancements",
               "Performance, Reporting & Advanced Analytics")

card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(3.7), Inches(5.0), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(3.3), Inches(0.4),
             "⚡ High-Speed Reporting", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Replaced openpyxl with Rust-based calamine engine",
    "• 5-10x faster Excel parsing and generation",
    "• Monthly Category & Brand Performance Report (APR vs MAY)",
    "• Values formatted in Crores (₹)",
    "• PDF + Excel dual export",
], Inches(1.95), left=Inches(0.8), width=Inches(3.3), font_size=11)

card2 = add_rounded_rect(slide, Inches(4.65), Inches(1.5), Inches(3.7), Inches(5.0), DARK_GRAY)
add_text_box(slide, Inches(4.85), Inches(1.55), Inches(3.3), Inches(0.4),
             "📋 New Dashboard Sections", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Monthly Retention Analysis",
    "• Campaign Analysis / Dormant Resurrection",
    "• Redemption Analysis (loyalty points)",
    "• Loyalty Point Cohort (year-over-year)",
    "• Customer Segmentation Download",
    "• Enterprise AI Chat Agent",
    "• n8n Chat Widget Integration",
], Inches(1.95), left=Inches(4.85), width=Inches(3.3), font_size=11)

card3 = add_rounded_rect(slide, Inches(8.7), Inches(1.5), Inches(3.7), Inches(5.0), DARK_GRAY)
add_text_box(slide, Inches(8.9), Inches(1.55), Inches(3.3), Inches(0.4),
             "📊 Advanced Metrics", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• DataTables: searchable & sortable grids",
    "• Dynamic ASP mapping",
    "• Product-to-category cross-referencing",
    "• Branch-level customer analytics",
    "• Multi-branch selection & custom date range",
    "• Role-based download access (mygadmin)",
    "• 5,242,619 unique customers verified",
], Inches(1.95), left=Inches(8.9), width=Inches(3.3), font_size=11)


# ══════════════════════════════════════════════════════════════
# SLIDE 9: BIGG BOSS & FoneFlix
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "7. BIGG BOSS & FoneFlix Registration Portals",
               "Flask-Based Registration + Video Upload Platform")

card1 = add_rounded_rect(slide, Inches(0.6), Inches(1.5), Inches(5.8), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(1.55), Inches(5), Inches(0.4),
             "🎥 Video Upload Architecture (Google Drive OAuth)", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Custom OAuth 2.0 Client ID integration for Google Drive uploads",
    "• Users upload large video files (up to 500MB) directly to Drive",
    "• Backup mechanism: Render server storage if Drive upload fails",
    "• Upload progress bar with real-time status feedback",
    "• JSON backup file for data redundancy",
], Inches(1.95), left=Inches(0.9), width=Inches(5.3), font_size=12)

card2 = add_rounded_rect(slide, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(7.0), Inches(1.55), Inches(5), Inches(0.4),
             "🎨 Dynamic UI/UX Design", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• OTT reality-show aesthetic with 45/60 split layout",
    "• Neon-glow typography, glassmorphism forms, 3D particle animations",
    "• Pixel-perfect responsive design (desktop + mobile)",
    "• Welcome popup banner with auto-display on load",
    "• Security: Server-side validation, 1 mobile = 1 entry",
], Inches(1.95), left=Inches(7.1), width=Inches(5.3), font_size=12)

card3 = add_rounded_rect(slide, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.5), DARK_GRAY)
add_text_box(slide, Inches(0.8), Inches(4.35), Inches(5), Inches(0.4),
             "🔄 Project Pivot & Rebranding: Bigg Boss → myG FoneFlix", font_size=14, color=GOLD, bold=True)
add_bullet_slide_content(slide, [
    "• Successfully transitioned entire codebase from Bigg Boss (Navy Blue/Neon Pink) to FoneFlix (Dark Browns/Orange) with zero downtime",
    "• Refactored: Theme colors, form text, consent clauses, submit button, upload labels, page title, meta info",
    "• New Google Drive folder and Google Sheets for FoneFlix contest data",
    "• Audition Closed page created and deployed when registration period ended",
    "• Deployed on Render with custom domain: biggboss.myg.in",
], Inches(4.75), left=Inches(0.9), width=Inches(11.5), font_size=12)


# ══════════════════════════════════════════════════════════════
# SLIDE 10: Projects Delivered Summary Table
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "All Projects Delivered — Q2 2026")

projects = [
    ("1", "OSG-myG-PORTAL",           "Claims Management + WhatsApp Automation",   "✅ Production", "Render"),
    ("2", "myG Loyalty Dashboard",     "Enterprise Analytics + AI Agent",           "✅ Production", "Render"),
    ("3", "SHE START Dashboard",       "Startup Evaluation System",                 "✅ Production", "Render"),
    ("4", "BIGBOSS → FoneFlix",        "Registration + Video Upload Portal",        "✅ Production", "Render"),
    ("5", "HR E-Signature Portal",     "Employee Signature System",                 "✅ Production", "Render"),
    ("6", "FIFA World Cup Contest",    "Prediction & Leaderboard Platform",         "✅ Production", "Vercel"),
    ("7", "Enterprise AI Agent",       "NLP → SQL Business Intelligence",           "✅ Production", "Integrated"),
    ("8", "n8n Chat Agent",            "Automated AI Chat Widget",                  "✅ Production", "n8n Cloud"),
]

# Table header
header_top = Inches(1.3)
add_shape(slide, Inches(0.6), header_top, Inches(11.7), Inches(0.5), GOLD)
cols = [Inches(0.6), Inches(1.2), Inches(4.0), Inches(7.5), Inches(9.6)]
headers = ["#", "Project", "Description", "Status", "Platform"]
widths_h = [Inches(0.5), Inches(2.7), Inches(3.4), Inches(2.0), Inches(2.0)]
for i, (col, hdr, w) in enumerate(zip(cols, headers, widths_h)):
    add_text_box(slide, col, header_top + Inches(0.05), w, Inches(0.4),
                 hdr, font_size=12, color=DARK_NAVY, bold=True, alignment=PP_ALIGN.LEFT)

# Table rows
for idx, (num, name, desc, status, platform) in enumerate(projects):
    row_top = header_top + Inches(0.55) + Inches(idx * 0.55)
    row_color = DARK_GRAY if idx % 2 == 0 else SLATE
    add_shape(slide, Inches(0.6), row_top, Inches(11.7), Inches(0.5), row_color)
    vals = [num, name, desc, status, platform]
    for col, val, w in zip(cols, vals, widths_h):
        c = GOLD if val.startswith("✅") else (WHITE if col == cols[1] else LIGHT)
        b = True if col == cols[1] else False
        add_text_box(slide, col, row_top + Inches(0.05), w, Inches(0.4),
                     val, font_size=11, color=c, bold=b)


# ══════════════════════════════════════════════════════════════
# SLIDE 11: Key Achievements
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.06), H, GOLD)
section_header(slide, "Key Achievements — Q2 2026")

achievements = [
    ("8 Production Apps",       "Built & deployed 8 production applications in a single quarter"),
    ("12.6M+ Rows",             "Managed massive database with 26 materialized views and sub-second performance"),
    ("AI/ML Models Live",       "Integrated Random Forest, LSTM, and LLM models into production dashboards"),
    ("WhatsApp Automation",     "Automated messaging for 4 claim status types using Telinfy API"),
    ("Real-time Sync",          "Built Google Sheets ↔ PostgreSQL sync across 3 different projects"),
    ("AI Speed: 5min → <1s",   "Reduced AI Agent response from 5-10 minutes to under 1 second"),
    ("OAuth Video Pipeline",    "Engineered OAuth 2.0 upload handling files up to 500MB"),
    ("Kerala AI Features",      "Built Malayalam Calendar featurizer for Kerala-specific predictive models"),
    ("Pixel-Perfect Design",    "Delivered responsive designs for desktop & mobile across all portals"),
    ("Zero-Downtime Pivot",     "Pivoted Bigg Boss portal to FoneFlix contest with zero downtime"),
]

for i, (title, desc) in enumerate(achievements):
    row = i // 2
    col = i % 2
    left = Inches(0.6) + col * Inches(6.2)
    top = Inches(1.3) + row * Inches(1.1)

    card = add_rounded_rect(slide, left, top, Inches(5.8), Inches(0.95), DARK_GRAY)
    # Number badge
    badge = add_rounded_rect(slide, left + Inches(0.15), top + Inches(0.15), Inches(0.5), Inches(0.5), GOLD)
    add_text_box(slide, left + Inches(0.15), top + Inches(0.15), Inches(0.5), Inches(0.5),
                 str(i+1), font_size=16, color=DARK_NAVY, bold=True, alignment=PP_ALIGN.CENTER)
    # Title
    add_text_box(slide, left + Inches(0.8), top + Inches(0.1), Inches(4.8), Inches(0.35),
                 title, font_size=13, color=GOLD, bold=True)
    # Description
    add_text_box(slide, left + Inches(0.8), top + Inches(0.45), Inches(4.8), Inches(0.45),
                 desc, font_size=11, color=LIGHT)


# ══════════════════════════════════════════════════════════════
# SLIDE 12: Thank You
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_NAVY)
add_shape(slide, Inches(0), Inches(0), Inches(0.08), H, GOLD)
add_shape(slide, Inches(0.08), Inches(0), W, Inches(0.02), GOLD)
add_shape(slide, Inches(0.08), H - Inches(0.02), W, Inches(0.02), GOLD)

add_text_box(slide, Inches(0), Inches(2.2), W, Inches(1.0),
             "Thank You", font_size=52, color=GOLD, bold=True,
             alignment=PP_ALIGN.CENTER, font_name="Calibri Light")

add_gold_line(slide, Inches(5.5), Inches(3.4), Inches(2.3), Inches(0.04))

add_text_box(slide, Inches(0), Inches(3.8), W, Inches(0.6),
             "Q2 2026 — April | May | June", font_size=20, color=WHITE,
             alignment=PP_ALIGN.CENTER, font_name="Calibri")

add_text_box(slide, Inches(0), Inches(4.5), W, Inches(0.5),
             "Jasil N Balussery  ·  Technology / Software Development  ·  myG",
             font_size=14, color=MED_GRAY, alignment=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════
output_path = r"c:\Users\jasil_myg\Desktop\OSG-myG-PORTAL-mainnnnn - Copy\Q2_2026_Work_Summary_Business.pptx"
prs.save(output_path)
print(f"\nPresentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
